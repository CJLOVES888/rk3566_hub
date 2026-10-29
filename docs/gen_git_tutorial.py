# -*- coding: utf-8 -*-
"""Generate beginner Git + GitHub Word tutorial."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime
import os

out_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "Git与GitHub代码管理基础教程.docx",
)

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)


def set_run_font(run, name_cn="宋体", name_en="Times New Roman", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name_cn)


def add_heading_cn(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        if level == 1:
            set_run_font(run, "黑体", "Arial", 16, True, RGBColor(0x1A, 0x56, 0x8C))
        elif level == 2:
            set_run_font(run, "黑体", "Arial", 14, True, RGBColor(0x2E, 0x75, 0xB6))
        else:
            set_run_font(run, "黑体", "Arial", 12, True, RGBColor(0x3B, 0x3B, 0x3B))
    return h


def add_para(text, size=12, bold=False, first_line_indent=True, align="justify"):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    if first_line_indent:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, "宋体", "Times New Roman", size, bold)
    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.left_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, "宋体", "Times New Roman", 12)
    return p


def add_code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, "Consolas", "Consolas", 10.5, False, RGBColor(0x20, 0x20, 0x20))
    return p


def set_cell_text(cell, text, bold=False, size=10.5, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, "宋体", "Times New Roman", size, bold)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, center=True)
        shade_cell(table.rows[0].cells[i], "D6E3F0")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], str(val))
    doc.add_paragraph()
    return table


# ========== Cover ==========
for _ in range(2):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Git 与 GitHub")
set_run_font(run, "黑体", "Arial", 22, True, RGBColor(0x1A, 0x56, 0x8C))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("代码管理与云端保存基础教程")
set_run_font(run, "黑体", "Arial", 18, True)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("面向初学者｜结合 RK3566 Hub 实际上传经验")
set_run_font(run, "宋体", "Times New Roman", 12)

doc.add_paragraph()
for k, v in [
    ("适用环境", "Windows PowerShell + Git + GitHub"),
    ("配套工程", "rk3566_hub（摄像头 / MQTT / SQLite）"),
    ("文档日期", datetime.date.today().strftime("%Y年%m月%d日")),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("%s：%s" % (k, v))
    set_run_font(run, "宋体", "Times New Roman", 12)

doc.add_page_break()

# ========== TOC ==========
add_heading_cn("目录", 1)
for t in [
    "1 为什么要用 Git 和 GitHub",
    "2 核心概念（先懂这几个词）",
    "3 安装与首次配置",
    "4 第一次把本地项目上传到 GitHub",
    "5 日常改代码：add → commit → push",
    "6 从云端拉回最新代码（pull）",
    "7 代理、登录与 Token（国内常见问题）",
    "8 .gitignore：哪些文件不要上传",
    "9 常用命令速查表",
    "10 安全注意与最佳实践",
    "附录 A 结合本项目的推荐操作清单",
    "附录 B 常见报错对照",
]:
    add_para(t, first_line_indent=False, align="left")

doc.add_page_break()

# ========== 1 ==========
add_heading_cn("1 为什么要用 Git 和 GitHub", 1)
add_para(
    "Git 是安装在电脑上的「版本管理工具」：每次重要改动都可以拍一张快照（commit），"
    "以后能回看、对比、回退，避免改坏代码找不到上一版。"
)
add_para(
    "GitHub 是网上的代码托管网站（同类还有 Gitee、GitLab）。把本地 Git 仓库"
    "push 上去后，代码就有了云端备份，换电脑、写简历、给面试官看仓库链接都更方便。"
)
add_para("一句话：Git 管历史，GitHub 管存放与分享。", bold=True, first_line_indent=False)

# ========== 2 ==========
add_heading_cn("2 核心概念（先懂这几个词）", 1)
add_table(
    ["名词", "通俗理解"],
    [
        ["仓库 Repository", "一个项目的全部文件 + 历史记录"],
        ["工作区 Working Tree", "你正在编辑的文件夹内容"],
        ["暂存区 Stage", "git add 之后、准备提交的清单"],
        ["提交 Commit", "一次带说明的快照，形成历史"],
        ["分支 Branch", "平行的开发线；新手常用 main"],
        ["远程 Remote", "云端仓库地址，一般叫 origin"],
        ["推送 Push", "把本地提交上传到 GitHub"],
        ["拉取 Pull", "把 GitHub 上的新提交下载合并到本地"],
        ["克隆 Clone", "第一次从网上把整个仓库下载到本地"],
    ],
)

# ========== 3 ==========
add_heading_cn("3 安装与首次配置", 1)
add_heading_cn("3.1 安装 Git", 2)
add_para("访问 https://git-scm.com/download/win 安装 Git for Windows。安装后新开 PowerShell：")
add_code("git --version")

add_heading_cn("3.2 设置姓名和邮箱", 2)
add_para("这两个信息会写进每次 commit，建议与 GitHub 账号邮箱一致：")
add_code('git config --global user.name "你的名字"\ngit config --global user.email "你的邮箱@example.com"')

add_heading_cn("3.3 查看配置", 2)
add_code("git config --global --list")

# ========== 4 ==========
add_heading_cn("4 第一次把本地项目上传到 GitHub", 1)
add_para(
    "本节对应你已经完成的 rk3566_hub 上传流程。若以后有新项目，按同样步骤即可。"
)

add_heading_cn("4.1 在 GitHub 网站新建空仓库", 2)
add_bullet("打开 github.com → 登录 → New repository（创建新的仓库）")
add_bullet("填写仓库名，例如 rk3566_hub")
add_bullet("可见度选 Public（公开作品）或 Private（私有）")
add_bullet("重要：不要勾选「Add a README / Add .gitignore」——本地已有代码时保持空仓库更简单")
add_bullet("不要点「Import」（导入是从别的网站搬家，不是上传本机代码）")
add_bullet("点击 Create repository，记下地址，例如 https://github.com/用户名/rk3566_hub.git")

add_heading_cn("4.2 在本地进入项目目录", 2)
add_code("cd E:\\xulian\\vscode\\RK3566_Project\\rk3566_hub")

add_heading_cn("4.3 初始化并做第一次提交", 2)
add_code(
    "git init\n"
    "# 先写好 .gitignore（见第 8 章），再：\n"
    "git add .\n"
    "git status\n"
    'git commit -m "feat: initial commit"'
)
add_para(
    "git status 时务必检查：没有私钥、密码、build 目录、数据库文件。确认无误再 commit。"
)

add_heading_cn("4.4 关联远程并推送", 2)
add_code(
    "git branch -M main\n"
    "git remote add origin https://github.com/你的用户名/rk3566_hub.git\n"
    "git push -u origin main"
)
add_para(
    "成功后终端常出现 set up to track 'origin/main'；之后在同一分支可直接 git push。"
    "若提示 Everything up-to-date，表示云端已与本地最新提交一致（你当前状态就是这样）。"
)

# ========== 5 ==========
add_heading_cn("5 日常改代码：add → commit → push", 1)
add_para("这是每天用得最多的三步循环：")
add_code(
    "cd E:\\xulian\\vscode\\RK3566_Project\\rk3566_hub\n"
    "\n"
    "# 1) 看改了哪些文件\n"
    "git status\n"
    "git diff\n"
    "\n"
    "# 2) 把改动放进暂存区\n"
    "git add .\n"
    "# 或只添加某个文件：\n"
    "# git add src/net/mqtt_client.cpp\n"
    "\n"
    "# 3) 写说明并生成快照\n"
    'git commit -m "fix: improve camera GStreamer pipeline"\n'
    "\n"
    "# 4) 上传到 GitHub\n"
    "git push"
)

add_heading_cn("5.1 提交说明怎么写", 2)
add_para("建议简短、说「为什么/做了什么」，可用前缀：")
add_table(
    ["前缀", "含义", "示例"],
    [
        ["feat", "新功能", "feat: add MQTT screen_off command"],
        ["fix", "修 bug", "fix: correct Bemfa clientId auth"],
        ["docs", "文档", "docs: add stage technical report"],
        ["refactor", "重构", "refactor: cleanup camera open logic"],
        ["chore", "杂务", "chore: update .gitignore"],
    ],
)

# ========== 6 ==========
add_heading_cn("6 从云端拉回最新代码（pull）", 1)
add_para(
    "如果你在另一台电脑改过并 push 了，或网页上改了文件，本地需要同步："
)
add_code("git pull")
add_para(
    "新手阶段尽量避免多人同时改同一文件。若 pull 时出现冲突（conflict），"
    "Git 会在文件里标出冲突块，需要手动保留正确内容后，再 git add 与 git commit。"
)

add_heading_cn("6.1 换电脑：克隆已有仓库", 2)
add_code("git clone https://github.com/你的用户名/rk3566_hub.git\ncd rk3566_hub")

# ========== 7 ==========
add_heading_cn("7 代理、登录与 Token（国内常见问题）", 1)
add_heading_cn("7.1 连不上 github.com", 2)
add_para("若出现 Failed to connect to github.com:443，多半是网络问题。可先测试：")
add_code("curl.exe -I https://github.com")
add_para(
    "若你使用 Clash 等代理，注意 Git 不会自动用「系统代理端口」。"
    "需把 Git 代理设成与软件一致的本地端口。例如代理显示 127.0.0.1:7897："
)
add_code(
    "git config --global http.proxy http://127.0.0.1:7897\n"
    "git config --global https.proxy http://127.0.0.1:7897"
)
add_para("端口必须和代理软件一致。曾误设为 7890 而实际是 7897 时，会出现 over proxy 失败。")
add_para("不用代理时取消：")
add_code(
    "git config --global --unset http.proxy\n"
    "git config --global --unset https.proxy"
)

add_heading_cn("7.2 GitHub 不再接受「账户密码」推送", 2)
add_para(
    "HTTPS 推送时，Password 位置应填写 Personal Access Token（PAT），不是网页登录密码。"
)
add_bullet("打开 https://github.com/settings/tokens")
add_bullet("Generate new token (classic) → 勾选 repo → 生成并复制")
add_bullet("git push 时 Username 填 GitHub 用户名，Password 粘贴 Token")

add_heading_cn("7.3 浏览器回调 127.0.0.1 拒绝连接", 2)
add_para(
    "有时 Git Credential Manager 会打开浏览器授权，授权成功后跳回 http://127.0.0.1:某端口，"
    "若本地监听已关闭，就会出现「无法访问此页面 / 拒绝连接」。此时不必纠结该页面，"
    "改用手动创建 Token + 在终端输入用户名和 Token 即可。"
)

add_heading_cn("7.4 403 Permission denied", 2)
add_para(
    "网络已通但仍 403，通常是 Token 权限不足、Token 过期，或 Windows 里缓存了旧密码。"
    "重新生成带 repo 权限的 Token，删除旧凭据后再 push。"
)

# ========== 8 ==========
add_heading_cn("8 .gitignore：哪些文件不要上传", 1)
add_para("在仓库根目录创建 .gitignore，例如：")
add_code(
    "build/\n"
    "**/build/\n"
    "*.o\n"
    "*.user\n"
    "CMakeCache.txt\n"
    "CMakeFiles/\n"
    "data/*.db\n"
    "*.db\n"
    ".env\n"
    "secrets/\n"
    ".vscode/\n"
    ".idea/\n"
    "Thumbs.db"
)
add_para(
    "绝对不要提交：巴法云私钥 UID、appID、secretKey、真实密码、内网机器密码。"
    "密钥用环境变量（如 HUB_MQTT_UID）在运行时注入。"
)

# ========== 9 ==========
add_heading_cn("9 常用命令速查表", 1)
add_table(
    ["命令", "作用"],
    [
        ["git status", "看当前改动与分支状态"],
        ["git diff", "看未暂存的具体差异"],
        ["git add <文件>", "加入暂存区"],
        ["git add .", "暂存当前目录下所有改动"],
        ["git commit -m \"说明\"", "生成一次提交"],
        ["git log --oneline", "简明查看提交历史"],
        ["git remote -v", "查看远程地址"],
        ["git push", "上传到远程"],
        ["git push -u origin main", "首次推送并设置上游跟踪"],
        ["git pull", "下载并合并远程更新"],
        ["git clone <url>", "克隆远程仓库"],
        ["git branch", "查看本地分支"],
        ["git checkout -b 新分支", "创建并切换到新分支（进阶）"],
    ],
)

# ========== 10 ==========
add_heading_cn("10 安全注意与最佳实践", 1)
add_bullet("先 git status，再 commit；先确认无密钥，再 push。")
add_bullet("小步提交：一个功能或一个 bug 一次 commit，方便回滚。")
add_bullet("main 保持可编译/可运行；大实验可另开分支（学会后再用）。")
add_bullet("Token 泄露后立刻在 GitHub 上撤销（Revoke）并换新。")
add_bullet("公开仓库默认任何人可读，简历项目注意脱敏。")
add_bullet("大文件（模型 .rknn、视频）慎传；可用 Release 或网盘，必要时用 Git LFS。")

# ========== Appendix A ==========
add_heading_cn("附录 A 结合本项目的推荐操作清单", 1)
add_bullet("仓库地址示例：https://github.com/CJLOVES888/rk3566_hub")
add_bullet("本地目录：E:\\xulian\\vscode\\RK3566_Project\\rk3566_hub")
add_bullet("已上传内容应包含：src/、scripts/、docs/ 技术文档、CMakeLists.txt 等")
add_bullet("日常改完相机/MQTT 后：git add . → git commit -m \"...\" → git push")
add_bullet("代理端口以 Clash「系统代理设置」为准（本机曾为 7897）")
add_bullet("运行板端程序仍用 export HUB_MQTT_UID=...，不要把 UID 写进仓库")

# ========== Appendix B ==========
add_heading_cn("附录 B 常见报错对照", 1)
add_table(
    ["报错摘要", "原因", "处理"],
    [
        ["Could not connect to server :443", "网络/未设代理或端口错", "检查代理；设对 http.proxy 端口"],
        ["over proxy 127.0.0.1 after 0 ms", "代理端口错误或代理未开", "改成实际端口如 7897；打开 Clash"],
        ["403 Permission denied", "Token/权限/旧凭据", "新建 repo 权限 Token，清凭据重推"],
        ["Authentication failed", "用户名或 Token 错", "Username=GitHub 名；Password=Token"],
        ["Everything up-to-date", "云端已是最新", "正常，无需再推；有新改动先 commit"],
        ["rejected (fetch first)", "远程有你本地没有的提交", "先 git pull 再 push"],
        ["not a git repository", "当前目录未 init", "cd 到项目根目录；必要时 git init"],
        ["127.0.0.1 拒绝连接（浏览器）", "OAuth 本地回调失败", "改用 Token 手动登录推送"],
    ],
)

doc.add_paragraph()
add_para(
    "—— 文档结束。掌握 add / commit / push / pull 四步，即可完成绝大多数个人项目云端保存。",
    first_line_indent=False,
    align="center",
)

doc.save(out_path)
print("OK", out_path)
print("size", os.path.getsize(out_path))
