#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
RK3566 智能值守终端 —— 面试知识链路与实现过程手册
对应简历「项目二」五条 bullet，按源码真实实现整理。
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = Path(__file__).resolve().parent / "RK3566值守终端_面试知识链路与实现手册.docx"


def set_run_font(run, size=11, bold=False, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_cn(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        set_run_font(p.add_run(text), size=16, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        set_run_font(p.add_run(text), size=13, bold=True, color=RGBColor(0x2E, 0x6D, 0xA4))
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    else:
        set_run_font(p.add_run(text), size=11.5, bold=True, color=RGBColor(0x33, 0x33, 0x33))
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    return p


def add_para(doc, text, *, bold=False, size=10.5):
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), size=size, bold=bold)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_run_font(p.add_run(text), size=10.5)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.4)
    return p


def add_qa(doc, q, a):
    add_para(doc, "Q：" + q, bold=True, size=10.5)
    add_para(doc, "A：" + a, size=10.5)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        title.add_run("RK3566 智能值守终端\n面试知识链路与实现过程手册"),
        size=18,
        bold=True,
        color=RGBColor(0x1F, 0x4E, 0x79),
    )
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        sub.add_run("对应简历项目二 · 按仓库真实代码整理 · 可直接背诵/口述"),
        size=10,
        color=RGBColor(0x66, 0x66, 0x66),
    )

    # ========== 0 总览 ==========
    add_heading_cn(doc, "0 一分钟项目口述稿（先背这段）", 1)
    add_para(
        doc,
        "这是一个跑在瑞芯微 RK3566 上的嵌入式 Linux 值守终端。我独立完成整机："
        "用 Qt5 做全屏 UI（无桌面走 linuxfb，有桌面走 xcb）；"
        "针对 Rockchip CSI 的 multiplanar 特点写了原生 V4L2 采帧；"
        "用帧差做有人亮屏 / 无人息屏；用 libmosquitto 对接巴法云，手机可远程控屏，板子上报事件；"
        "本地用 SQLite 记审计日志。NPU 检测线程和 YOLO 后处理接口已经接进工程，"
        "等 .rknn 模型到位就能做人形告警。技术栈是 C++17 + Qt + V4L2 + MQTT + SQLite + CMake。",
    )

    add_heading_cn(doc, "0.1 系统数据流（知识主链路）", 2)
    add_code(
        doc,
        "摄像头 /dev/video-camera0\n"
        "  → V4l2MplaneCapture（NV12→BGR）\n"
        "  → CameraWorker 线程 emit frameReady(QImage)\n"
        "       ├→ MainWindow::onFrame → 预览 / 息屏叠层\n"
        "       ├→ PresenceGuard::onFrame → person_near / person_away → MQTT + 亮灭屏\n"
        "       └→ Detector::onFrame（有模型时）→ person_detected → MQTT\n"
        "手机/网页 --MQTT--> rk3566cmd → MqttClient → screen_on/off → UI\n"
        "板端事件 --MQTT--> rk3566event；同时 SQLite event_log 落库",
    )
    add_para(
        doc,
        "面试官若问「模块怎么拆」，答：视觉采集、值守算法、网络、存储、UI 五层；"
        "UI 只做编排与展示，重活放在工作线程 / mosquitto 内部 loop。",
    )

    add_heading_cn(doc, "0.2 关键源码地图", 2)
    add_bullet(doc, "入口：src/main.cpp → QApplication + MainWindow 全屏")
    add_bullet(doc, "UI 编排：src/ui/main_window.cpp（接相机/值守/MQTT/DB/Detector）")
    add_bullet(doc, "采帧：src/vision/camera_worker.cpp + v4l2_mplane_capture.cpp")
    add_bullet(doc, "值守：src/vision/presence_guard.cpp")
    add_bullet(doc, "MQTT：src/net/mqtt_client.cpp + protocol.h")
    add_bullet(doc, "存储：src/storage/database.cpp")
    add_bullet(doc, "NPU 预留：src/vision/detector.cpp + rknn_engine.c")
    add_bullet(doc, "启动：scripts/run.sh（显示后端 / 环境变量）")

    # ========== 1 相机 ==========
    add_heading_cn(doc, "1 相机采集：为何必须做 V4L2 MPLANE", 1)
    add_heading_cn(doc, "1.1 问题背景（面试高频）", 2)
    add_para(
        doc,
        "Rockchip CSI 摄像头在 Linux 上通常暴露为 V4L2 multiplanar 设备。"
        "OpenCV 直接 CAP_V4L2 打开 /dev/video0 常失败或得到「假彩色 / 单色帧」"
        "（格式协商不对，把 NV12 当 RGB 解）。"
        "所以项目优先走原生 ioctl 路径，OpenCV/GStreamer 仅作回退。",
    )
    add_heading_cn(doc, "1.2 实现过程（对应简历第 1 条）", 2)
    add_bullet(doc, "打开设备：open(O_RDWR|O_NONBLOCK) → VIDIOC_QUERYCAP 判断 CAPTURE_MPLANE")
    add_bullet(doc, "设格式：VIDIOC_S_FMT，像素格式 NV12（可环境变量切 NV21），分辨率 640×480")
    add_bullet(doc, "缓冲：VIDIOC_REQBUFS + mmap 多缓冲；QUEUE/DQBUF 循环取帧")
    add_bullet(doc, "色彩：NV12 平面拼紧凑缓冲 → OpenCV cvtColor 转 BGR → 再转 RGB 填 QImage")
    add_bullet(doc, "线程：CameraWorker 在独立 QThread 里 while(m_running) 读帧，emit frameReady(img.copy())")
    add_bullet(doc, "设备选择：HUB_CAMERA_DEVICE → /dev/video-camera0 → 扫描 video0..15")
    add_heading_cn(doc, "1.3 你要能讲清的概念", 2)
    add_bullet(doc, "V4L2：Video4Linux2，Linux 视频设备统一接口（ioctl 命令集）")
    add_bullet(doc, "MPLANE：一帧可分多 plane（Y / UV）；RK CSI 常见 1 plane 里塞 NV12")
    add_bullet(doc, "NV12：YUV 420 半平面，Y 全分辨率 + UV 交错，比 RGB 省带宽")
    add_bullet(doc, "mmap：把内核驱动缓冲映射到用户态，零拷贝/少拷贝取帧")
    add_bullet(doc, "为何 copy()：跨线程传递 QImage 需保证生命周期，简化内存所有权（可优化点）")
    add_heading_cn(doc, "1.4 面试问答", 2)
    add_qa(
        doc,
        "为什么不用纯 OpenCV？",
        "RK CSI 是 multiplanar，OpenCV 对格式协商不友好，曾出现打不开或颜色错误；"
        "原生 V4L2 可控 pixelformat/stride，颜色正确后再转给 Qt。",
    )
    add_qa(
        doc,
        "取帧卡死怎么办？",
        "NONBLOCK + select/poll 或读失败 sleep 重试；工程上可再做连续失败后 reopen 设备。",
    )

    # ========== 2 值守 ==========
    add_heading_cn(doc, "2 值守：PresenceGuard 帧差算法", 1)
    add_heading_cn(doc, "2.1 设计目标", 2)
    add_para(
        doc,
        "不依赖 NPU 也能做「门口有人 / 无人」粗判断：有运动→亮屏并上报；"
        "长时间无显著运动→息屏。纯 Qt QImage 实现，不绑 OpenCV，CPU 占用低。",
    )
    add_heading_cn(doc, "2.2 实现过程（对应简历第 2 条）", 2)
    add_bullet(doc, "隔帧：processEvery（默认 2），不是每帧都算")
    add_bullet(doc, "预处理：转灰度 + 缩到 previewW×previewH（如 160×120），FastTransformation")
    add_bullet(doc, "差分：逐像素 |cur-prev| ≥ pixelDiff 计为变化点；score = changed/total")
    add_bullet(doc, "判定运动：score ≥ motionRatio（默认约 0.025，即 2.5% 像素变化）")
    add_bullet(doc, "去抖：连续 confirmHits（默认 3）次运动才从「无人」切到「有人」→ emit presenceNear")
    add_bullet(doc, "息屏：已占用状态下，距上次运动超过 idleMs（默认 30s）→ emit presenceAway")
    add_bullet(doc, "MainWindow：near → setScreenOn(true)+MQTT；away → setScreenOn(false)+MQTT；防抖 3s")
    add_heading_cn(doc, "2.3 参数从哪来", 2)
    add_para(
        doc,
        "环境变量 HUB_PRESENCE / HUB_PRESENCE_IDLE_SEC / HUB_PRESENCE_THRESH / HUB_PRESENCE_CONFIRM，"
        "也可从 SQLite config 表读默认值。",
    )
    add_heading_cn(doc, "2.4 面试问答", 2)
    add_qa(
        doc,
        "帧差和 YOLO 人物检测有什么区别？",
        "帧差只认「画面是否在动」，窗帘/光线也会误触发，但算力极低，适合息屏策略；"
        "YOLO 认「是不是人」，适合告警。项目里两者分工：帧差管亮灭屏，NPU 管确认人形。",
    )
    add_qa(
        doc,
        "如何降低误触发？",
        "提高 motionRatio、增加 confirmHits、加大缩图滤波、拉长 idle；必要时加 ROI 只盯门口区域。",
    )

    # ========== 3 MQTT ==========
    add_heading_cn(doc, "3 MQTT：巴法云双向控制", 1)
    add_heading_cn(doc, "3.1 协议与主题", 2)
    add_bullet(doc, "Broker：mqtt.bemfa.com:9501")
    add_bullet(doc, "认证方式一：ClientId = 用户私钥 UID（用户名密码可空）")
    add_bullet(doc, "下行主题 rk3566cmd：手机/网页 → 板子（板子 subscribe）")
    add_bullet(doc, "上行主题 rk3566event：板子 → 云端（publish 事件/ACK）")
    add_bullet(
        doc,
        "JSON 信封（protocol.h）："
        '{"type":"cmd|event|ack","id":"...","ts":...,"payload":{...}}',
    )
    add_heading_cn(doc, "3.2 实现过程（对应简历第 3 条）", 2)
    add_bullet(doc, "mosquitto_lib_init → mosquitto_new(clientId) → 注册 connect/disconnect/message 回调")
    add_bullet(doc, "mosquitto_connect_async + mosquitto_loop_start（库内网络线程）")
    add_bullet(doc, "reconnect_delay_set(2,30,true)：断线指数退避重连")
    add_bullet(doc, "CONNACK 成功后 subscribe cmdTopic QoS1")
    add_bullet(
        doc,
        "回调在 mosquitto 线程 → QMetaObject::invokeMethod(..., QueuedConnection) 回到 Qt 对象线程，保证线程安全",
    )
    add_bullet(doc, "MainWindow 解析 payload.action：screen_on / screen_off；执行后 publish ACK")
    add_bullet(doc, "值守/测试按钮/NPU 事件同样走 publishEvent → rk3566event")
    add_heading_cn(doc, "3.3 踩过的坑（面试加分）", 2)
    add_bullet(doc, "Lookup error：DNS 失败（板子没配网/resolv.conf），不是 UID 错")
    add_bullet(doc, "手机与板端都用同一 UID 作 ClientId 会互踢")
    add_bullet(doc, "第三方客户端发指令有时要 topic/set；网页控制台可直接发 screen_off")
    add_heading_cn(doc, "3.4 面试问答", 2)
    add_qa(
        doc,
        "为什么选 MQTT 而不是自己做 TCP？",
        "IoT 场景需要穿透与主题订阅，巴法云已提供设备云；QoS/心跳/重连由库和 Broker 承担，"
        "应用层专注 JSON 业务。点餐项目已练过裸 TCP，本项目练云端对接。",
    )
    add_qa(
        doc,
        "QoS 怎么选？",
        "控屏与关键事件用 QoS1（至少一次）；可接受重复，靠业务 id/防抖处理。",
    )

    # ========== 4 架构 ==========
    add_heading_cn(doc, "4 Qt 多线程架构与显示后端", 1)
    add_heading_cn(doc, "4.1 线程模型（对应简历第 4 条）", 2)
    add_bullet(doc, "UI 线程：QApplication 事件循环，MainWindow 刷新控件")
    add_bullet(doc, "相机线程：CameraWorker::moveToThread，started → start() 采帧循环")
    add_bullet(doc, "检测线程：Detector（有模型时）同样 QThread，跳帧 + busy 丢帧防堆积")
    add_bullet(doc, "MQTT：libmosquitto loop 线程；经 QueuedConnection 投递到 MqttClient/UI")
    add_bullet(doc, "跨线程通信：信号槽（默认队列连接），避免在相机线程直接改 UI")
    add_heading_cn(doc, "4.2 linuxfb vs xcb（必考）", 2)
    add_para(
        doc,
        "Qt 通过 QT_QPA_PLATFORM 选平台插件：linuxfb 直接写 /dev/fb0，适合无桌面产品形态；"
        "xcb 连 X 服务器，适合有 XFCE 的调试。桌面还在时用 linuxfb 会被挡住；"
        "关掉桌面后若仍留着 DISPLAY=:0，run.sh 会误判走 xcb 导致 could not connect to display。"
        "正确：全屏前 unset DISPLAY，或 HUB_FORCE_LINUXFB=1。",
    )
    add_heading_cn(doc, "4.3 MainWindow 编排顺序", 2)
    add_code(
        doc,
        "initStorage() → buildUi() → startPresence() → startCamera()\n"
        "→ startDetector() → startMqtt() → 时钟定时器 + 清理 7 天前事件",
    )
    add_heading_cn(doc, "4.4 面试问答", 2)
    add_qa(
        doc,
        "信号槽为什么能跨线程？",
        "元对象系统把槽调用排入接收者线程事件队列（QueuedConnection），"
        "参数类型需注册（如 QVector<DetectBox>）。",
    )
    add_qa(
        doc,
        "如何保证 UI 不卡？",
        "重 I/O 不放 UI 线程；预览可限帧；值守缩图隔帧；NPU busy 时丢帧。",
    )

    # ========== 5 存储与 NPU ==========
    add_heading_cn(doc, "5 SQLite 审计与 NPU 预留", 1)
    add_heading_cn(doc, "5.1 SQLite（对应简历第 5 条前半）", 2)
    add_bullet(doc, "路径：build/bin/data/smart_hub.db")
    add_bullet(doc, "PRAGMA journal_mode=WAL；synchronous=NORMAL，兼顾嵌入式写入性能")
    add_bullet(doc, "表 event_log：类型/详情/来源（mqtt/camera/presence/npu/ui）")
    add_bullet(doc, "表 config：mqtt_host、主题、presence 参数、device_id 等")
    add_bullet(doc, "MQTT 断线时事件仍 logEvent，保证可追溯")
    add_heading_cn(doc, "5.2 NPU 预留（对应简历第 5 条后半，诚实讲）", 2)
    add_para(
        doc,
        "当前进度：Detector + rknn_engine（letterbox/推理/NMS）代码就绪；"
        "CMake 可选链接 librknnrt（HUB_HAVE_RKNN）；"
        "MainWindow 已 connect personDetected → 亮屏 + SQLite + MQTT。"
        "缺：板端 models/yolov8n.rknn（需 PC Linux 用 rknn-toolkit2 从 ONNX 转换）。"
        "面试不要说「已经完成 YOLO 告警落地」，要说「链路打通，待模型联调」。",
    )
    add_heading_cn(doc, "5.3 面试问答", 2)
    add_qa(
        doc,
        "为什么用 SQLite 不用直接写文件？",
        "结构化查询、配置键值、并发读写有 WAL；比自制日志格式好维护。",
    )

    # ========== 6 构建运行 ==========
    add_heading_cn(doc, "6 构建、部署与运行（体现工程能力）", 1)
    add_bullet(doc, "主机改代码 → scp 到板 → 板端 cmake && make（本地 aarch64 编译）")
    add_bullet(doc, "依赖：Qt5、OpenCV、libmosquitto、SQLite 插件")
    add_code(
        doc,
        "# 有桌面调试\n"
        "export HUB_MQTT_UID=私钥; export DISPLAY=:0; unset HUB_QT_PLATFORM\n"
        "bash scripts/run.sh\n\n"
        "# 无桌面全屏\n"
        "sudo systemctl isolate multi-user.target\n"
        "unset DISPLAY; export HUB_FORCE_LINUXFB=1; export HUB_QT_PLATFORM=linuxfb\n"
        "bash scripts/run.sh",
    )

    # ========== 7 深挖题 ==========
    add_heading_cn(doc, "7 可能的深挖题速查", 1)
    add_qa(doc, "C++17 用了什么？", "QThread+atomic 标志位、STL vector 缓冲、智能指针思想（Qt 父子所有权）、结构化配置。")
    add_qa(doc, "和点餐 LVGL 项目对比？", "LVGL 偏轻量控件+自绘framebuffer；本项目 Qt 生态+多线程+CSI/MQTT 系统集成，更接近产品中控。")
    add_qa(doc, "最大难点？", "CSI multiplanar 颜色/节点选择；linuxfb 与桌面抢屏；巴法云 ClientId/DNS 联调。")
    add_qa(doc, "若入职做 GUI 优化，你会先做什么？", "UI 限帧与 FastTransformation；run.sh 探测 DISPLAY 是否真可用；相机失败自动 reopen。")
    add_qa(doc, "安全上注意什么？", "UID 只放环境变量不进仓库；MQTT 公网注意主题权限；设备物理部署注意隐私（摄像头）。")

    # ========== 8 背诵卡片 ==========
    add_heading_cn(doc, "8 五条简历 bullet → 30 秒展开卡", 1)
    add_para(doc, "① V4L2：RK CSI 要用 MPLANE，NV12 自转 BGR，OpenCV 作回退。", bold=True)
    add_para(doc, "② 值守：缩图帧差 + 连续命中 + idle 超时，驱动亮灭屏和 MQTT 事件。", bold=True)
    add_para(doc, "③ MQTT：巴法云 UID 作 ClientId，cmd 订、event 发，控屏+ACK，回调回 Qt 线程。", bold=True)
    add_para(doc, "④ 架构：相机/UI/检测分线程，信号槽解耦；linuxfb/xcb 由 run.sh 选型。", bold=True)
    add_para(doc, "⑤ SQLite WAL 审计；RKNN/Detector 已接线，等 rknn 模型做人形告警。", bold=True)

    add_para(
        doc,
        "建议用法：先背第 0 章口述稿 + 第 8 章五条；再按面试官追问翻到对应章节。"
        "对着板子跑一遍「亮屏/息屏/网页 screen_off」，把现象和日志关键词记下，比只背理论更稳。",
        size=10,
    )

    doc.save(str(OUT))
    print("OK", OUT)


if __name__ == "__main__":
    main()
