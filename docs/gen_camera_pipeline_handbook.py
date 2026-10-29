#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""摄像头与画面处理 —— 面试深水区专项手册（对照 rk3566_hub 源码）"""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = Path(__file__).resolve().parent / "摄像头与画面处理_深水区专项手册.docx"


def font(run, size=11, bold=False, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def h(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        font(p.add_run(text), 15, True, RGBColor(0x1F, 0x4E, 0x79))
    elif level == 2:
        font(p.add_run(text), 12.5, True, RGBColor(0x2E, 0x6D, 0xA4))
    else:
        font(p.add_run(text), 11, True, RGBColor(0x33, 0x33, 0x33))
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)


def p(doc, text, size=10.5, bold=False):
    para = doc.add_paragraph()
    font(para.add_run(text), size, bold)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.25


def b(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    font(para.add_run(text), 10.5)
    para.paragraph_format.space_after = Pt(1)


def code(doc, text):
    para = doc.add_paragraph()
    r = para.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(9)
    para.paragraph_format.left_indent = Cm(0.3)
    para.paragraph_format.space_after = Pt(5)


def qa(doc, q, a):
    p(doc, "Q：" + q, bold=True)
    p(doc, "A：" + a)


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(2.0)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(t.add_run("摄像头采集 × 画面处理\n深水区专项手册"), 17, True, RGBColor(0x1F, 0x4E, 0x79))
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(s.add_run("对照 V4l2MplaneCapture / CameraWorker / PresenceGuard / MainWindow::onFrame"), 9.5, False, RGBColor(0x66, 0x66, 0x66))

    h(doc, "写在前面：你要建立的心智模型", 1)
    p(doc, "摄像头不是「直接给你一张彩色 JPG」，而是硬件不断往内核缓冲里塞「原始像素」。"
          "软件要做的事只有三步：①协商格式并取走缓冲；②把原始格式变成程序能用的颜色空间；"
          "③分发给预览 / 值守 /（将来）NPU。你项目把①②放在相机线程，③用信号槽分发。")
    code(doc,
         "传感器 → ISP/驱动 → /dev/video*（NV12 原始帧）\n"
         "  → ioctl DQBUF 取出 mmap 缓冲\n"
         "  → NV12 打包 → cvtColor → BGR → RGB → QImage\n"
         "  → emit frameReady\n"
         "       ├ PresenceGuard（灰度缩图帧差）\n"
         "       ├ MainWindow 预览（缩放 setPixmap）\n"
         "       └ Detector（可选，以后）")

    # ========== 基础概念 ==========
    h(doc, "一、先补 6 个基础概念（不懂后面全糊）", 1)

    h(doc, "1.1 设备节点 /dev/video*", 2)
    p(doc, "Linux 把摄像头抽象成字符设备。一块板上可能有很多 videoN："
          "有的是真实采集口，有的是 ISP/编码中间节点。Rockchip CSI 常用符号链接 "
          "/dev/video-camera0 指向真正的 capture 节点。所以代码优先用它，而不是盲开 video0。")

    h(doc, "1.2 V4L2 = Video4Linux2", 2)
    p(doc, "一套 ioctl 命令：查能力、设宽高/像素格式、申请缓冲、开流、取帧、停流。"
          "你不用记全部命令号，但要能按顺序背：QUERYCAP → S_FMT → REQBUFS → QUERYBUF+mmap "
          "→ QBUF → STREAMON →（循环 DQBUF→处理→QBUF）→ STREAMOFF。")

    h(doc, "1.3 single-planar vs multiplanar（MPLANE）——本项目核心坑", 2)
    p(doc, "单平面：一帧一个内存块。多平面：一帧可拆成多个 plane（例如 Y 一块、UV 一块）。"
          "Rockchip CSI 驱动大量用 CAPTURE_MPLANE。OpenCV 老路径常按「单平面」去谈，"
          "结果打开失败，或谈成了格式但按错布局解，出现整屏发绿/发粉/像噪声。"
          "因此项目写了 V4l2MplaneCapture，直接跟驱动用 MPLANE 协议对话。")

    h(doc, "1.4 像素格式：RGB vs YUV / NV12", 2)
    b(doc, "RGB888：每像素 R,G,B 各 1 字节，Qt 显示友好，占带宽大。")
    b(doc, "YUV：亮度和色度分离。人眼对亮度更敏感，色度可抽样，省带宽。")
    b(doc, "NV12（YUV420 半平面）：先整幅 Y（宽×高），再半高的 UV 交错行（宽×高/2）。总大小 = 宽×高×1.5。")
    b(doc, "NV21：和 NV12 类似，但 U/V 字节顺序对调。环境变量 HUB_CAMERA_YUV=nv21 可切。")
    p(doc, "摄像头输出 NV12 很常见；显示器/Qt 要 RGB。所以必须有一次色彩转换。")

    h(doc, "1.5 stride（步长）≠ width", 2)
    p(doc, "一行像素在内存里可能有对齐填充。width=640 时 stride 可能仍是 640，"
          "也可能是 704 等。若按 width 去读，画面会「斜纹撕裂」。"
          "代码里 yuvToBgr 按 stride 跳行拷贝，再拼成紧凑 NV12 再转色。")

    h(doc, "1.6 mmap 环形缓冲", 2)
    p(doc, "驱动在内核准备若干缓冲（本项目申请约 3 个）。用户态 mmap 映射进来。"
          "QBUF=把空缓冲交给驱动去装下一帧；DQBUF=取回已装满的缓冲。"
          "取完必须再 QBUF 还回去，否则缓冲耗尽，采帧停住。这就是「环形队列」思想。")

    # ========== 采集实现 ==========
    h(doc, "二、采集实现：V4l2MplaneCapture 逐步对照", 1)
    p(doc, "文件：src/vision/v4l2_mplane_capture.cpp", bold=True)

    h(doc, "2.1 open()：把设备和格式谈妥", 2)
    b(doc, "open(device, O_RDWR|O_NONBLOCK)：非阻塞，配合后面 select，避免永久卡死。")
    b(doc, "VIDIOC_QUERYCAP：确认有 VIDEO_CAPTURE_MPLANE 或普通 CAPTURE。")
    b(doc, "VIDIOC_S_FMT：宽高 640×480，pixelformat=NV12，MPLANE 下先试 1 plane，失败再试 2 planes。")
    b(doc, "记下驱动真正返回的 width/height/stride（驱动可能改你申请的值）。")
    b(doc, "initMmap + queueAll + STREAMON：缓冲就绪并开始出流。")

    h(doc, "2.2 initMmap / queueAll", 2)
    b(doc, "REQBUFS count=3，MEMORY_MMAP。")
    b(doc, "对每个 index：QUERYBUF 拿长度和偏移 → mmap → 得到 start[] 指针。")
    b(doc, "全部 QBUF 一次，驱动才有空槽可写。")

    h(doc, "2.3 readBgr()：取一帧并转成 BGR Mat", 2)
    b(doc, "select(fd, 2s超时)：等到有数据可读。")
    b(doc, "VIDIOC_DQBUF：拿到某个 buf.index。")
    b(doc, "若双 plane：start[0]=Y，start[1]=UV，各自带 stride，送进 yuvToBgr。")
    b(doc, "若单 plane：Y 与 UV 首尾相接；stride==width 可直接 cvtColor，否则仍走打包。")
    b(doc, "处理完立刻 VIDIOC_QBUF 归还。")
    b(doc, "输出：cv::Mat，类型 BGR，连续内存，三通道。")

    h(doc, "2.4 yuvToBgr：为什么要「先打包」", 2)
    p(doc, "OpenCV 的 COLOR_YUV2BGR_NV12 期望输入是紧凑布局（每行有效宽=width）。"
          "驱动缓冲常有 stride 填充，所以先按行 memcpy 成宽×高×1.5 的 packed，再 cvtColor。"
          "这是正确性优先；性能上多一次拷贝，640×480 可接受。")

    # ========== CameraWorker ==========
    h(doc, "三、CameraWorker：谁在循环、如何发出去", 1)
    p(doc, "文件：src/vision/camera_worker.cpp；运行在独立 QThread。", bold=True)

    h(doc, "3.1 设备候选顺序", 2)
    code(doc, "HUB_CAMERA_DEVICE → /dev/video-camera0 →（无环境变量时）扫描 /dev/video0..15")

    h(doc, "3.2 路径 A（当前实测成功路径）", 2)
    p(doc, "对每个候选设备试 V4l2MplaneCapture::open；成功则进入 while(m_running)：")
    code(doc,
         "native.readBgr(&bgr)\n"
         "cv::cvtColor(bgr, rgb, COLOR_BGR2RGB)   // OpenCV 默认 BGR，Qt 要 RGB\n"
         "QImage img(rgb.data, w, h, rgb.step, Format_RGB888)\n"
         "emit frameReady(img.copy())            // copy：跨线程必须自有一份像素\n"
         "统计 FPS → emit fpsChanged")
    p(doc, "失败 sleep(30) 继续，避免空转打满 CPU。路径 A 成功则直接 return，不走路径 B。")

    h(doc, "3.3 路径 B（回退）", 2)
    p(doc, "原生失败时，用 OpenCV VideoCapture + 多种 GStreamer v4l2src 管道试探。"
          "CSI 上曾出现「假成功」：管道通了但颜色仍错，所以优先保证路径 A。"
          "面试可说：策略是原生优先、管线兜底。")

    h(doc, "3.4 为什么必须 img.copy()", 2)
    p(doc, "QImage(rgb.data, ...) 是「借用」Mat 的内存。循环下一圈 rgb 被覆盖/析构后，"
          "若信号槽异步排队还拿着旧指针就会花屏/崩溃。copy() 深拷贝一份像素交给接收方。")

    # ========== 画面处理 ==========
    h(doc, "四、画面处理：帧到 UI / 值守各走哪条路", 1)

    h(doc, "4.1 总分流（MainWindow::onFrame）", 2)
    code(doc,
         "frameReady → UI 线程 onFrame(frame)\n"
         "  1) presence->onFrame(frame)     // 息屏也要喂，否则无法亮屏\n"
         "  2) if (!screenOn) return;       // 不刷新预览，省 CPU\n"
         "  3) drawBoxes(frame)             // 有检测框才画\n"
         "  4) scaled → setPixmap           // 适配 QLabel 大小")

    h(doc, "4.2 预览路径（你看见的画面）", 2)
    b(doc, "输入：Format_RGB888 的 QImage（与采集端一致）。")
    b(doc, "drawBoxes：用 QPainter 画矩形和标签；无框则原图返回。")
    b(doc, "QPixmap::fromImage + scaled(KeepAspectRatio, SmoothTransformation)。")
    p(doc, "注意：Smooth 缩放质量好但更吃 CPU；面试可提「可改 FastTransformation + 限帧」作优化。")

    h(doc, "4.3 值守路径（PresenceGuard）——另一类「画面处理」", 2)
    p(doc, "目的不是好看，而是便宜地判断「有没有动」。步骤：")
    b(doc, "隔帧：processEvery（默认每 2 帧算 1 次）")
    b(doc, "转灰度 Grayscale8（丢掉色度，运动主要看亮度）")
    b(doc, "缩到 160×120（像素少 → 算得快）")
    b(doc, "与上一帧逐像素 |Δ| ≥ 25 计变化；score = 变化像素 / 总像素")
    b(doc, "score≥0.025 算运动；连续 3 次运动才 presenceNear；30s 无运动 presenceAway")
    p(doc, "这和「显示」是两条线：显示要保真彩色；值守要降维提特征。面试强调「同帧多消费者、处理目标不同」。")

    h(doc, "4.4 颜色空间对照表（背下来）", 2)
    code(doc,
         "驱动输出     NV12/YUV\n"
         "OpenCV 中间  BGR Mat     （cvtColor NV12→BGR，再 BGR→RGB）\n"
         "Qt 预览      RGB888 QImage\n"
         "值守内部     Gray8 缩图\n"
         "（将来 NPU） 多数要 RGB/BGR letterbox 到 640×640")

    # ========== 线程与性能 ==========
    h(doc, "五、线程、拷贝与性能（深挖常问）", 1)
    b(doc, "采集在 CameraWorker 线程；UI/值守在主线程（值守目前由 onFrame 同步调用）。")
    b(doc, "每帧至少：NV12→BGR、BGR→RGB、QImage copy；预览再缩放。这是 CPU 热点。")
    b(doc, "可优化方向：UI 限 15fps；Fast 缩放；值守直接收更小图；共享缓冲减 copy（复杂度高）。")
    b(doc, "息屏跳过 setPixmap，但仍跑 PresenceGuard——产品正确性优先。")

    # ========== 排障 ==========
    h(doc, "六、排障清单（结合你真实经历）", 1)
    b(doc, "打不开：节点选错（不是 camera0）；权限；被其他进程占用。→ probe_camera.sh / v4l2-ctl")
    b(doc, "能开但花屏/单色：当 RGB 解了 NV12，或 stride 没处理。→ 必须走 MPLANE+正确 cvtColor")
    b(doc, "绿/粉偏色：NV12/NV21 搞反。→ HUB_CAMERA_YUV=nv21 试")
    b(doc, "卡顿：Smooth 缩放+全帧 copy；或 UI 与采集同频。→ 限帧/Fast")
    b(doc, "select 超时：传感器没出流、接线/驱动问题、STREAMON 失败")

    # ========== 口述 ==========
    h(doc, "七、面试 90 秒口述（摄像头专项）", 1)
    p(doc, "RK 的 CSI 在 Linux 上是 V4L2 multiplanar，输出多是 NV12。"
          "我不用 OpenCV 硬开，而是自己用 ioctl 做 QUERYCAP、S_FMT、mmap 环形缓冲，"
          "DQBUF 取帧后按 stride 拼成紧凑 NV12，再转成 BGR/RGB 交给 Qt。"
          "CameraWorker 在独立线程循环取帧，frameReady 把 QImage 拷贝发出去。"
          "UI 线程一边给 PresenceGuard 做灰度缩图帧差，一边在亮屏时缩放显示；"
          "息屏时只做值守不算预览。这样显示正确、占用可控，也避开了 CSI 上 OpenCV 的坑。")

    h(doc, "八、高频问答", 1)
    qa(doc, "NV12 和 RGB 差在哪？",
       "NV12 是 YUV420 半平面，体积约为 RGB 的一半；采集端省带宽，显示前再转 RGB。")
    qa(doc, "什么是 stride？",
       "一行字节跨度，可能大于宽度；忽略它会导致图像错位。")
    qa(doc, "DQBUF / QBUF 干什么？",
       "取出已填充缓冲 / 归还空缓冲，形成环形队列。")
    qa(doc, "为什么要 QImage::copy？",
       "跨线程异步传递时，必须拥有独立像素，不能借用下一帧会覆盖的缓冲区。")
    qa(doc, "值守为什么转灰度还缩小？",
       "运动主要体现在亮度变化；缩小后计算量按像素数下降，嵌入式上更合适。")
    qa(doc, "和 GStreamer 比呢？",
       "GStreamer 管道灵活但在 CSI 上曾出现假成功；原生 V4L2 对格式更可控，故作首选。")

    h(doc, "九、对照源码阅读顺序（建议 1 小时）", 1)
    b(doc, "v4l2_mplane_capture.h/.cpp：只看 open → initMmap → readBgr → yuvToBgr")
    b(doc, "camera_worker.cpp：candidateDevices + 路径 A 的 while 循环")
    b(doc, "main_window.cpp：onFrame / drawBoxes")
    b(doc, "presence_guard.cpp：灰度、缩放、差分、状态机")
    p(doc, "读的时候在纸上画「这一帧的内存形态从 NV12 变成了什么」，比只看函数名有效。")

    doc.save(str(OUT))
    print("OK", OUT)


if __name__ == "__main__":
    main()
