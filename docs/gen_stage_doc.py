# -*- coding: utf-8 -*-
"""Generate RK3566 hub stage technical Word document."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime
import os

out_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "RK3566智能中控终端_阶段技术文档.docx",
)

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)


def set_run_font(run, name_cn="宋体", name_en="Times New Roman", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name_cn)


def add_heading_cn(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        if level == 1:
            set_run_font(run, "黑体", "Arial", 16, True, RGBColor(0x1A, 0x56, 0x8C))
        elif level == 2:
            set_run_font(run, "黑体", "Arial", 14, True, RGBColor(0x2E, 0x75, 0xB6))
        else:
            set_run_font(run, "黑体", "Arial", 12, True, RGBColor(0x3B, 0x3B, 0x3B))
    return h


def add_para(text, size=12, bold=False, first_line_indent=True, align="justify"):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    if first_line_indent:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, "宋体", "Times New Roman", size, bold)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.left_indent = Cm(0.74 + level * 0.5)
    run = p.add_run(text)
    set_run_font(run, "宋体", "Times New Roman", 12)
    return p


def set_cell_text(cell, text, bold=False, size=10.5, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, "宋体", "Times New Roman", size, bold)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, center=True)
        shade_cell(table.rows[0].cells[i], "D6E3F0")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], str(val), center=False)
    doc.add_paragraph()
    return table


# ===================== Cover =====================
for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("RK3566 边缘 AI 智能安防中控终端")
set_run_font(run, "黑体", "Arial", 22, True, RGBColor(0x1A, 0x56, 0x8C))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("（rk3566_hub）")
set_run_font(run, "黑体", "Arial", 16, True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("阶段技术文档")
set_run_font(run, "黑体", "Arial", 18, True)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(
    "当前阶段：摄像头预览 + MQTT 双向遥控打通\n"
    "NPU 推理链路已预留（待 .rknn 模型）"
)
set_run_font(run, "宋体", "Times New Roman", 12)

doc.add_paragraph()
doc.add_paragraph()

meta = [
    ("项目名称", "RK3566 Hub 智能中控终端"),
    ("工程版本", "v0.3"),
    ("硬件平台", "KickPi / Rockchip RK3566"),
    ("文档日期", datetime.date.today().strftime("%Y年%m月%d日")),
    ("文档性质", "阶段完成情况与技术点说明"),
]
for k, v in meta:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("%s：%s" % (k, v))
    set_run_font(run, "宋体", "Times New Roman", 12)

doc.add_page_break()

# ===================== TOC =====================
add_heading_cn("目录", 1)
toc_items = [
    "1 项目概述",
    "2 当前阶段完成情况总览",
    "3 系统架构与模块划分",
    "4 已完成功能详述",
    "5 技术点详尽清单",
    "6 通信协议与消息规范",
    "7 构建、部署与运行",
    "8 联调验证记录与结论",
    "9 未完成与下一阶段计划",
    "10 目录结构与关键文件说明",
    "11 秋招项目表述建议（可选）",
    "附录 A 环境变量一览",
    "附录 B 常见问题与排障要点",
]
for t in toc_items:
    add_para(t, first_line_indent=False, align="left")

doc.add_page_break()

# ===================== 1 =====================
add_heading_cn("1 项目概述", 1)
add_heading_cn("1.1 项目定位", 2)
add_para(
    "本项目面向嵌入式 Linux / 边缘 AI 方向，基于 Rockchip RK3566 开发板构建一套"
    "「本地采集 + 边缘推理（规划）+ 云端/手机远程控制 + 本地持久化」的智能中控终端。"
    "工程名称为 rk3566_hub，当前以无图形桌面环境的 linuxfb 全屏 Qt 应用形式运行于板端，"
    "目标场景包括智能安防预览、远程关屏/开屏、事件上报与后续人物检测告警联动。"
)
add_heading_cn("1.2 硬件与运行环境", 2)
add_table(
    ["类别", "内容"],
    [
        ["SoC", "Rockchip RK3566（四核 Cortex-A55，含 NPU）"],
        ["板卡", "KickPi（板端用户 kickpi，局域网联调）"],
        ["显示", "Framebuffer /dev/fb0，分辨率实测 1024×600"],
        ["摄像头", "CSI 树莓派风格摄像头，节点 /dev/video-camera0"],
        ["OS", "板端 Debian/Ubuntu 类 Linux"],
        ["开发主机", "Windows + Cursor/VS Code，源码 scp 同步后板端编译"],
        ["云平台", "巴法云 MQTT 设备云（mqtt.bemfa.com:9501）"],
        ["手机端", "IoT MQTT Panel（指令下发 + 事件订阅）"],
    ],
)
add_heading_cn("1.3 本阶段目标与达成", 2)
add_para(
    "本阶段优先打通「可演示的完整闭环」：板端启动后能显示摄像头画面，"
    "手机/网页可远程控制屏幕开关，板端可向云端上报状态与 ACK。"
    "NPU/YOLOv8 推理代码与 CMake 可选链接已就绪，但运行时尚未加载 .rknn 模型"
    "（启动横幅显示 HUB_RKNN_MODEL=<missing>），故人物检测告警属于下一阶段。"
)

# ===================== 2 =====================
add_heading_cn("2 当前阶段完成情况总览", 1)
add_table(
    ["模块", "状态", "完成度说明"],
    [
        ["工程骨架 / CMake", "已完成", "C++17，Qt5 AUTOMOC，可选 RKNN"],
        ["主界面 UI（linuxfb）", "已完成", "全屏预览、状态栏、时钟、测试上报按钮"],
        ["摄像头采集", "已完成", "CSI + GStreamer BGR，640×480 稳定出帧"],
        ["MQTT 客户端", "已完成", "libmosquitto 异步连接/订阅/发布/重连"],
        ["巴法云对接", "已完成", "方式一 ClientId=私钥；主题 rk3566cmd/event"],
        ["手机遥控开关屏", "已完成", "IoT MQTT Panel + appID/secretKey + /set"],
        ["协议层 JSON", "已完成", "cmd/event/ack 统一封装，兼容纯文本指令"],
        ["SQLite 持久化", "已完成", "event_log + config，启动可写库"],
        ["NPU Detector 框架", "代码就绪", "QThread + rknn_engine 封装，缺模型"],
        ["YOLOv8 .rknn 推理", "未启用", "运行日志 HUB_RKNN_MODEL=<missing>"],
        ["人物检测 MQTT 告警", "待联调", "依赖 NPU 模型与 personDetected 链路"],
    ],
)
add_para(
    "结论：截至文档日期，项目已完成「显示 + 采集 + 远程控制 + 本地存储」主链路，"
    "MQTT 双向已实测通过；边缘 AI 推理处于「接口与工程集成完成、模型未接入」状态。"
)

# ===================== 3 =====================
add_heading_cn("3 系统架构与模块划分", 1)
add_heading_cn("3.1 逻辑架构", 2)
add_para(
    "系统采用分层模块化设计：表现层（MainWindow）负责 UI 与业务编排；"
    "视觉层（CameraWorker / Detector / rknn_engine）负责采集与推理；"
    "网络层（MqttClient + protocol）负责云端通信与消息规范化；"
    "存储层（Database）负责事件与配置持久化。摄像头采集与 NPU 检测分别运行在独立 QThread，"
    "通过 Qt 信号槽（QueuedConnection）与 UI 线程解耦，避免阻塞界面刷新。"
)
add_heading_cn("3.2 模块职责", 2)
add_table(
    ["模块", "主要源文件", "职责"],
    [
        ["入口", "src/main.cpp", "QApplication、全屏几何、进入事件循环"],
        ["UI 编排", "src/ui/main_window.*", "预览、开关屏、MQTT/NPU/SQLite 联动"],
        ["摄像头", "src/vision/camera_worker.*", "设备探测、GStreamer/V4L2 采帧、FPS"],
        ["检测器", "src/vision/detector.*", "跳帧推理、画框数据、人物冷却上报"],
        ["RKNN 引擎", "src/vision/rknn_engine.c/.h", "C API 封装 letterbox/推理/NMS"],
        ["MQTT", "src/net/mqtt_client.*", "连接、订阅 cmd、发布 event、日志回调"],
        ["协议", "src/net/protocol.h", "JSON 消息构建与解析"],
        ["数据库", "src/storage/database.*", "SQLite 单例、事件日志、键值配置"],
        ["脚本", "scripts/*.sh", "启动、摄像头探测、MQTT 冒烟、RKNN 安装"],
    ],
)
add_heading_cn("3.3 运行时数据流", 2)
add_bullet("摄像头线程：V4L2/GStreamer 采集 → QImage → frameReady → UI 显示（可叠加检测框）。")
add_bullet("检测线程（可选）：frameReady → Detector::onFrame → rknn_engine_detect → detectionsUpdated / personDetected。")
add_bullet("下行控制：手机/网页 → 巴法云 rk3566cmd(/set) → 板端订阅回调 → handleCommand(screen_on/off) → UI 关屏叠层。")
add_bullet("上行事件：板端 publishEvent → rk3566event → 手机文本控件 / 网页消息区；本地同时 logEvent 写入 SQLite。")

# ===================== 4 =====================
add_heading_cn("4 已完成功能详述", 1)

add_heading_cn("4.1 板端全屏 UI（linuxfb）", 2)
add_para(
    "应用通过 QT_QPA_PLATFORM=linuxfb 直接渲染到帧缓冲，无需 X11/Wayland。"
    "主界面包含标题、视频预览区、关屏遮罩、MQTT 状态、NPU 状态、时钟与「上报测试事件」按钮。"
    "收到 screen_off 时切换到遮罩页（屏幕关闭语义）；screen_on 恢复预览。"
)

add_heading_cn("4.2 CSI 摄像头预览", 2)
add_para(
    "RK3566 CSI 摄像头多为 multiplanar V4L2，OpenCV 单纯 CAP_V4L2 打开 /dev/video0 常失败。"
    "CameraWorker 采用多级策略：环境变量 HUB_CAMERA_DEVICE → /dev/video-camera0 → 扫描 video0..15 → "
    "失败则走 GStreamer v4l2src 管道输出 BGR。当前实测路径为 /dev/video-camera0 (gst)，分辨率 640×480，三通道 BGR。"
)
add_bullet("技术点：V4L2 multiplanar、GStreamer 管道、OpenCV VideoCapture、QThread 采帧、FPS 统计。")
add_bullet("已知现象：OpenCV GStreamer Cannot query video position 警告对实时流可忽略，不影响预览。")

add_heading_cn("4.3 MQTT 双向通信（巴法云）", 2)
add_para(
    "板端使用 libmosquitto：mosquitto_connect_async + mosquitto_loop_start 后台网络循环；"
    "连接成功后订阅 rk3566cmd（QoS1）；事件与 ACK 发布到 rk3566event。"
    "认证采用巴法云官方「方式一」：ClientId = 用户私钥 UID，用户名/密码为空。"
    "断线后由库内指数退避重连，onConnect 中重新 subscribe。"
)
add_para(
    "手机端 IoT MQTT Panel 不能与板端共用同一 ClientId（否则互踢）。"
    "采用「方式二」：ClientId=phone01，Username=appID，Password=secretKey"
    "（通过 POST /vs/web/v1/userSecretKey 生成）。下发指令主题需使用 rk3566cmd/set，"
    "否则可能只更新云端显示而不推送到订阅设备——这是巴法云 MQTT 的关键约定。"
)

add_heading_cn("4.4 远程开关屏与 ACK", 2)
add_para(
    "板端 onMqttMessage 同时兼容：纯文本载荷（screen_on / screen_off）与标准 JSON。"
    "执行指令后通过 publishEvent 上报 ack（含 device_id、for、fps、screen_on 等字段），"
    "并写入 SQLite。网页控制台直接发送与手机按钮发送均可验证。"
)

add_heading_cn("4.5 SQLite 本地持久化", 2)
add_para(
    "Database 单例基于 QtSql，库文件位于 build/bin/data/smart_hub.db。"
    "表 event_log 记录摄像头/MQTT/系统/UI 事件；表 config 保存 mqtt_host、主题名等默认配置，"
    "并支持旧默认值迁移（如 bemfa.com → mqtt.bemfa.com）。"
)

add_heading_cn("4.6 NPU 检测框架（代码完成、模型未接入）", 2)
add_para(
    "Detector 运行于独立线程，默认每 2 帧推理 1 次，人物事件带冷却（约 2s）防刷屏。"
    "rknn_engine 以 C 接口封装：加载 .rknn、RGB letterbox、NPU 推理、YOLOv8 后处理与 NMS。"
    "CMake 通过 HUB_HAVE_RKNN 控制是否链接 librknnrt；未找到运行库时以 stub 编译，保证工程可构建。"
    "当前运行未设置有效 HUB_RKNN_MODEL，故 NPU 状态为未启用。"
)

# ===================== 5 =====================
add_heading_cn("5 技术点详尽清单", 1)
add_heading_cn("5.1 嵌入式与系统", 2)
add_bullet("Rockchip RK3566 边缘计算平台与 NPU 资源规划")
add_bullet("Linux framebuffer（linuxfb）无桌面图形栈直出 UI")
add_bullet("输入设备：evdev 触摸参数预留（QT_QPA_EVDEV_TOUCHSCREEN_PARAMETERS）")
add_bullet("设备节点管理：/dev/fb0、/dev/video-camera0、CSI multiplanar")
add_bullet("板端交叉/本地编译工作流：主机编辑 + scp + 板端 cmake/make")

add_heading_cn("5.2 C++ / Qt 应用框架", 2)
add_bullet("C++17、C11 混编（UI/业务 C++，RKNN 引擎 C）")
add_bullet("Qt5：Core / Gui / Widgets / Sql")
add_bullet("CMake AUTOMOC / AUTOUIC / AUTORCC")
add_bullet("QThread + Worker 对象移动线程，信号槽 QueuedConnection 跨线程投递")
add_bullet("QStackedLayout 实现预览/关屏叠层切换")
add_bullet("QTimer 时钟刷新；QLabel / QPushButton 状态展示与手动上报")
add_bullet("QProcessEnvironment 读取运行配置；QJsonDocument 序列化")

add_heading_cn("5.3 计算机视觉与多媒体", 2)
add_bullet("OpenCV：core / imgproc / videoio")
add_bullet("V4L2 摄像头枚举与打开失败回退策略")
add_bullet("GStreamer 管道（v4l2src → BGR）适配 Rockchip CSI")
add_bullet("QImage 与 OpenCV Mat 帧格式衔接（BGR 640×480）")
add_bullet("实时 FPS 统计与 UI 显示")
add_bullet("检测框叠加绘制（drawBoxes，为 NPU 结果预留）")

add_heading_cn("5.4 边缘 AI（RKNN）", 2)
add_bullet("Rockchip RKNN Runtime（librknnrt）可选链接")
add_bullet("YOLOv8n .rknn 模型加载路径约定（models/ 或 HUB_RKNN_MODEL）")
add_bullet("Letterbox 预处理、置信度阈值、NMS 后处理")
add_bullet("COCO class_id=0（person）事件抽取与冷却防抖")
add_bullet("编译期宏 HUB_HAVE_RKNN 实现有/无 NPU 双模式")

add_heading_cn("5.5 网络与物联网", 2)
add_bullet("MQTT 3.1.1（巴法云兼容）")
add_bullet("libmosquitto：异步连接、独立 loop 线程、自动重连")
add_bullet("QoS1 指令/事件可靠投递偏好；Retain 默认关闭")
add_bullet("巴法云方式一：ClientId=私钥 UID")
add_bullet("巴法云方式二：ClientId 随意 + appID/secretKey（手机端）")
add_bullet("巴法云 /set 推送语义：第三方客户端需发布到 topic/set")
add_bullet("主题模型：rk3566cmd（下行）、rk3566event（上行）")
add_bullet("连接诊断：CONNACK 原因字符串、mosquitto log 回调、订阅成功日志")

add_heading_cn("5.6 数据与协议", 2)
add_bullet("统一 JSON 信封：type / id / ts / payload")
add_bullet("消息类型：cmd、data、event、ack")
add_bullet("兼容纯文本指令（便于 App 按钮快速配置）")
add_bullet("SQLite：事件日志、键值配置、默认值种子与迁移")
add_bullet("事件源字段：mqtt / ui / camera / npu / system")

add_heading_cn("5.7 工程化与运维脚本", 2)
add_bullet("CMake 工程组织与依赖探测（Qt/OpenCV/pkg-config mosquitto/RKNN）")
add_bullet("qmake 备选工程文件 rk3566_hub.pro")
add_bullet("run.sh：linuxfb 环境、UID/摄像头/模型检查、启动二进制")
add_bullet("probe_camera.sh：摄像头节点探测辅助")
add_bullet("mqtt_smoke_test.sh：mosquitto_sub 冒烟订阅")
add_bullet("install_rknn_runtime.sh：运行库安装辅助")

add_heading_cn("5.8 联调与工具链", 2)
add_bullet("巴法云 Web 控制台在线发消息与订阅者状态观察")
add_bullet("IoT MQTT Panel：连接编辑、按钮面板、文本事件面板")
add_bullet("PowerShell Invoke-RestMethod 调用巴法云 userSecretKey 接口")
add_bullet("串口/SSH 日志验收：connected / subscribed / MQTT 收包日志")

# ===================== 6 =====================
add_heading_cn("6 通信协议与消息规范", 1)
add_heading_cn("6.1 Broker 与主题", 2)
add_table(
    ["项", "值"],
    [
        ["Host", "mqtt.bemfa.com（或 bemfa.com）"],
        ["Port", "9501（TCP）；加密可选 9503"],
        ["板端 ClientId", "用户私钥 UID"],
        ["板端 Username/Password", "空"],
        ["手机 ClientId", "如 phone01（不可与板端相同）"],
        ["手机 Username/Password", "appID / secretKey"],
        ["下行主题（板订阅）", "rk3566cmd"],
        ["手机发布主题", "rk3566cmd/set"],
        ["上行主题（板发布）", "rk3566event"],
    ],
)
add_heading_cn("6.2 JSON 消息示例", 2)
add_para("标准信封（板上行 ACK / 事件）示例：", first_line_indent=False)
add_para(
    '{"type":"ack","id":"4a5448c9","ts":1793012745,'
    '"payload":{"action":"ack","device_id":"dev01","for":"screen_on","fps":15.4,"screen_on":true}}',
    first_line_indent=False,
    align="left",
)
add_para(
    "手机按钮常用纯文本载荷（板端已兼容）：screen_off / screen_on。",
    first_line_indent=False,
)

add_heading_cn("6.3 指令处理逻辑", 2)
add_bullet("解析成功且 payload.action 存在 → 按 action 处理")
add_bullet("非 JSON → 整段 trim 后当作 action")
add_bullet("screen_off / screen_on → 切换 UI 叠层并 ACK")
add_bullet("其它 action → 记录日志，便于后续扩展（如抓拍、灵敏度等）")

# ===================== 7 =====================
add_heading_cn("7 构建、部署与运行", 1)
add_heading_cn("7.1 依赖安装（板端示例）", 2)
add_para(
    "需要：构建工具链、CMake、Qt5 开发包、OpenCV、libmosquitto-dev、SQLite 插件等。"
    "RKNN 需将 librknnrt.so 安装到系统库路径并 ldconfig（脚本 install_rknn_runtime.sh 可辅助）。"
)
add_heading_cn("7.2 编译", 2)
add_para(
    "cd ~/rk3566_hub && mkdir -p build && cd build && cmake .. && make -j$(nproc)",
    first_line_indent=False,
    align="left",
)
add_heading_cn("7.3 运行", 2)
add_para(
    "export HUB_MQTT_UID=<私钥> ; bash ~/rk3566_hub/scripts/run.sh",
    first_line_indent=False,
    align="left",
)
add_para(
    "可选环境变量：HUB_CAMERA_DEVICE、HUB_MQTT_HOST/PORT、HUB_MQTT_CMD_TOPIC、HUB_MQTT_EVENT_TOPIC、"
    "HUB_MQTT_CLIENT_ID、HUB_MQTT_USERNAME、HUB_MQTT_PASSWORD、HUB_RKNN_MODEL。"
)

# ===================== 8 =====================
add_heading_cn("8 联调验证记录与结论", 1)
add_heading_cn("8.1 验收现象（已通过）", 2)
add_bullet("启动日志出现 MQTT connected OK / MQTT subscribed: rk3566cmd")
add_bullet("摄像头 Camera OK GST /dev/video-camera0，首帧 640×480")
add_bullet("巴法云网页发送 screen_on/off 可控制板端屏幕")
add_bullet("IoT MQTT Panel 按钮发送后串口出现 MQTT 收包：rk3566cmd + screen_off/on")
add_bullet("网页 rk3566cmd 订阅者显示「在线 1」；事件区可见 ACK JSON")
add_heading_cn("8.2 关键排障经验（本阶段沉淀）", 2)
add_bullet("错误把板端 ClientId 改成 UID_后缀会导致巴法云鉴权失败，仅见 connecting 不见 subscribed")
add_bullet("手机与板端共用 UID 作 ClientId 会互踢，表现为半通/不稳定")
add_bullet("手机填 Username=私钥不属于方式二，需 appID/secretKey")
add_bullet("第三方客户端发布需用 topic/set，网页控制台发送不受此限制")
add_bullet("PowerShell 下 curl JSON 易转义失败，推荐 Invoke-RestMethod 生成密钥")

# ===================== 9 =====================
add_heading_cn("9 未完成与下一阶段计划", 1)
add_table(
    ["事项", "说明", "优先级"],
    [
        ["准备 yolov8n.rknn", "放入 models/ 或设置 HUB_RKNN_MODEL", "高"],
        ["确认 librknnrt", "板端 /usr/lib 可链接，CMake HUB_HAVE_RKNN=1", "高"],
        ["人物检测联调", "personDetected → UI/MQTT/SQLite 全链路验证", "高"],
        ["告警策略", "置信度阈值、冷却、夜间模式等可配置化", "中"],
        ["画框体验", "检测框叠加性能与刷新策略优化", "中"],
        ["安全加固", "密钥不入库明文、日志脱敏、主题权限说明", "中"],
        ["简历材料", "架构图、演示视频、指标（延迟/FPS/推理耗时）", "中"],
    ],
)

# ===================== 10 =====================
add_heading_cn("10 目录结构与关键文件说明", 1)
add_para("工程根目录 rk3566_hub 主要结构如下：", first_line_indent=False)
items = [
    ("CMakeLists.txt", "主构建脚本，声明依赖与 RKNN 可选开关"),
    ("rk3566_hub.pro", "qmake 备选工程"),
    ("src/main.cpp", "程序入口"),
    ("src/ui/", "主窗口 UI 与业务编排"),
    ("src/vision/", "摄像头、检测器、RKNN 引擎"),
    ("src/net/", "MQTT 客户端与协议头文件"),
    ("src/storage/", "SQLite 封装"),
    ("scripts/", "运行与探测脚本"),
    ("third_party/rknn/", "RKNN 头文件等第三方材料"),
    ("models/", "放置 .rknn 模型（约定路径）"),
    ("docs/", "技术文档输出目录（本文档）"),
]
for name, desc in items:
    add_bullet("%s — %s" % (name, desc))

# ===================== 11 =====================
add_heading_cn("11 秋招项目表述建议（可选）", 1)
add_para(
    "可在简历中强调：基于 RK3566 独立完成边缘端 Qt 应用；打通 CSI 摄像头在 multiplanar/"
    "GStreamer 场景下的稳定采集；使用 libmosquitto 对接公有 MQTT 云，完成手机 App 远程控制与"
    "事件回传；设计统一 JSON 协议与 SQLite 本地审计；预留 RKNN YOLOv8 推理线程模型，"
    "体现「物联网 + 嵌入式 + 边缘 AI」综合能力。量化指标建议补充：预览分辨率与 FPS、"
    "指令端到端时延、（下一阶段）NPU 单帧推理耗时。"
)

# ===================== Appendix =====================
add_heading_cn("附录 A 环境变量一览", 1)
add_table(
    ["变量", "作用", "示例/默认"],
    [
        ["HUB_MQTT_UID", "巴法云私钥（板端 ClientId）", "必填才能连云"],
        ["HUB_MQTT_HOST", "Broker 地址", "mqtt.bemfa.com"],
        ["HUB_MQTT_PORT", "端口", "9501"],
        ["HUB_MQTT_CMD_TOPIC", "指令主题", "rk3566cmd"],
        ["HUB_MQTT_EVENT_TOPIC", "事件主题", "rk3566event"],
        ["HUB_MQTT_CLIENT_ID", "覆盖 ClientId", "默认=UID"],
        ["HUB_MQTT_USERNAME", "方式二用户名", "默认空"],
        ["HUB_MQTT_PASSWORD", "方式二密码", "默认空"],
        ["HUB_CAMERA_DEVICE", "摄像头设备", "/dev/video-camera0"],
        ["HUB_RKNN_MODEL", ".rknn 路径", "models/yolov8n.rknn"],
        ["QT_QPA_PLATFORM", "Qt 平台插件", "linuxfb"],
        ["QT_QPA_FB", "帧缓冲设备", "/dev/fb0"],
    ],
)

add_heading_cn("附录 B 常见问题与排障要点", 1)
add_bullet("只有 MQTT connecting、无 subscribed：检查 ClientId 是否等于私钥；查网络与 9501 连通性。")
add_bullet("网页能控、App 不能控：检查手机 appID/secretKey；主题是否为 rk3566cmd/set。")
add_bullet("App 一连板子就掉线：ClientId 冲突，手机勿填私钥。")
add_bullet("无画面：确认 CSI 接好；运行 probe_camera.sh；检查 HUB_CAMERA_DEVICE。")
add_bullet("NPU 不工作：确认 librknnrt 与 .rknn；看启动横幅 HUB_RKNN_MODEL 是否 missing。")
add_bullet("SQLite 失败：确认 build/bin/data 可写。")

doc.add_paragraph()
add_para(
    "—— 文档结束。本文档描述截至当前阶段的真实完成情况，随 NPU 联调推进可继续增补版本记录。",
    first_line_indent=False,
    align="center",
)

doc.save(out_path)
print("OK", out_path)
print("size", os.path.getsize(out_path))
